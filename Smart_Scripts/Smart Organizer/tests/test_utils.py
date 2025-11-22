import unittest
from helpers.utils import extract_text_from_document
import shutil
from pathlib import Path
from fpdf import FPDF
from docx import Document
from openpyxl import Workbook
from pptx import Presentation


class TestDocumentTextExtraction(unittest.TestCase):
    
    def test_pdf_extraction(self):
      text = extract_text_from_document(self.PDF_FILE)
      self.assertIn("This is a sample PDF file for testing.", text)
    
    def test_docx_extraction(self):
      text = extract_text_from_document(self.DOCX_FILE)
      self.assertIn("This is a sample DOCX file for testing.", text)
    
    def test_txt_extraction(self):
      text = extract_text_from_document(self.TEXT_FILE)
      self.assertIn("This is a sample text file for testing.", text)
    
    def test_xlsx_extraction(self):
      text = extract_text_from_document(self.XLSX_FILE)
      self.assertIn("This", text)
      self.assertIn("is", text)
      self.assertIn("sample", text)
      self.assertIn("XLSX", text)
      self.assertIn("file", text)
      self.assertIn("for", text)
      self.assertIn("testing.", text)
    
    def test_pptx_extraction(self):
      text = extract_text_from_document(self.PPTX_FILE)
      self.assertIn("This is a sample PPTX file for testing.", text)
    
    def test_md_extraction(self):
      text = extract_text_from_document(self.MD_FILE)
      self.assertIn("This is a sample markdown file for testing.", text)
      
    def setUp(self):
      
      super().setUp()
      
      self.TEST_DIR = Path("temp_test_files")
      self.TEXT_FILE = self.TEST_DIR / "sample.txt"
      self.PDF_FILE = self.TEST_DIR / "sample.pdf"
      self.DOCX_FILE = self.TEST_DIR / "sample.docx"
      self.XLSX_FILE = self.TEST_DIR / "sample.xlsx"
      self.PPTX_FILE = self.TEST_DIR / "sample.pptx"
      self.MD_FILE = self.TEST_DIR / "sample.md"
      
      if not self.TEST_DIR.exists():
        self.TEST_DIR.mkdir(parents=True, exist_ok=True)
        
      if not self.PDF_FILE.exists():
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.cell(200, 10, txt="This is a sample PDF file for testing.", ln=True)
        pdf.output(self.PDF_FILE)
        
      if not self.DOCX_FILE.exists():
        doc = Document()
        doc.add_paragraph("This is a sample DOCX file for testing.")
        doc.save(self.DOCX_FILE)
      
      if not self.XLSX_FILE.exists():
        wb = Workbook()
        ws = wb.active
        ws.append(["This", "is", "a", "sample", "XLSX", "file", "for", "testing."])
        wb.save(self.XLSX_FILE)
      
      if not self.PPTX_FILE.exists():
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only layout
        title = slide.shapes.title
        title.text = "This is a sample PPTX file for testing."
        prs.save(self.PPTX_FILE)
      
      if not self.TEXT_FILE.exists():
        self.TEXT_FILE.write_text("This is a sample text file for testing.", encoding='utf-8')
        
      if not self.MD_FILE.exists():
        self.MD_FILE.write_text("# Sample Markdown\nThis is a sample markdown file for testing.", encoding='utf-8')
      
      
    def tearDown(self):
      
      super().tearDown()

      if self.TEST_DIR.exists():
          shutil.rmtree(self.TEST_DIR)
       